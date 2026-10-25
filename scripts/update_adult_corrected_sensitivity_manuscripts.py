"""Reconcile the main and supplementary manuscripts to corrected Adult sensitivities."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "adult_corrected_rerun"
AGE_NEW = RUN / "age_conditional_mar_outputs" / "age_conditional_mar_statistics.csv"
TYPE_NEW = RUN / "typeaware_imputation_outputs" / "typeaware_imputation_sensitivity_summary.csv"
MANUSCRIPT_IN = ROOT / "manuscript" / "FMG_manuscript_fullrun_reconciled.docx"
MANUSCRIPT_OUT = ROOT / "manuscript" / "FMG_manuscript_adult_sensitivities_reconciled.docx"
SUPP_IN = ROOT / "manuscript" / "FMG_supplementary_material.docx"
SUPP_OUT = ROOT / "manuscript" / "FMG_supplementary_adult_sensitivities_reconciled.docx"


def replace_prefix(doc, prefix, text):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}; got {len(matches)}")
    p = matches[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    p.add_run(text)


def write_cell(cell, text):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    p.add_run(str(text))
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def write_table(table, rows):
    if len(table.rows) != len(rows):
        raise ValueError(f"Table row mismatch: document={len(table.rows)}, new={len(rows)}")
    for row_obj, values in zip(table.rows, rows):
        if len(row_obj.cells) != len(values):
            raise ValueError("Table column mismatch")
        for cell, value in zip(row_obj.cells, values):
            write_cell(cell, value)


def replace_figure(doc, caption_prefix, image_path):
    matches = [p for p in doc.paragraphs if p.text.startswith(caption_prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one caption starting {caption_prefix!r}; got {len(matches)}")
    element = matches[0]._p.getprevious()
    while element is not None and not element.xpath('.//a:blip'):
        element = element.getprevious()
    if element is None:
        raise ValueError(f"No figure paragraph found before {caption_prefix!r}")
    blips = element.xpath('.//a:blip')
    if len(blips) != 1:
        raise ValueError(f"Expected one embedded image before {caption_prefix!r}; got {len(blips)}")
    rel_id, _ = doc.part.get_or_add_image(str(image_path))
    blips[0].set(qn('r:embed'), rel_id)


def r4(x):
    return f"{float(x):.4f}"


def load_sensitivities():
    comp = pd.read_csv(RUN / "complementary_group_summary_corrected_adult.csv")
    base = pd.read_csv(RUN / "baseline_recall_summary_corrected_adult.csv")
    group_mar = pd.read_csv(RUN / "group_conditional_mar_descriptive_summary_with_corrected_adult.csv")
    type_old = pd.read_csv(RUN / "typeaware_imputation_summary_with_corrected_adult.csv")
    type_new = pd.read_csv(TYPE_NEW)
    type_all = pd.concat([type_old[type_old.dataset != "Adult Income"], type_new], ignore_index=True)
    type_old_raw = pd.read_csv(ROOT / "sensitivity" / "typeaware_imputation_sensitivity_per_seed.csv")
    type_new_raw = pd.read_csv(RUN / "typeaware_imputation_outputs" / "typeaware_imputation_sensitivity_per_seed.csv")
    type_raw_all = pd.concat([type_old_raw[type_old_raw.dataset != "Adult Income"], type_new_raw], ignore_index=True)
    age_old = pd.read_csv(ROOT / "sensitivity" / "age_mar_strength_outputs" / "age_conditional_mar_statistics.csv")
    age_new = pd.read_csv(AGE_NEW)
    age_all = pd.concat([age_old[age_old.dataset != "Adult Income"], age_new], ignore_index=True)
    age_old_raw = pd.read_csv(ROOT / "sensitivity" / "age_mar_strength_outputs" / "age_conditional_mar_per_seed.csv")
    age_new_raw = pd.read_csv(RUN / "age_conditional_mar_outputs" / "age_conditional_mar_per_seed.csv")
    age_raw_all = pd.concat([age_old_raw[age_old_raw.dataset != "Adult Income"], age_new_raw], ignore_index=True)
    type_all.to_csv(RUN / "typeaware_imputation_summary_reconciled_latest.csv", index=False, float_format="%.8f")
    type_raw_all.to_csv(RUN / "typeaware_imputation_per_seed_reconciled_latest.csv", index=False, float_format="%.8f")
    age_all.to_csv(RUN / "age_conditional_mar_summary_reconciled_latest.csv", index=False, float_format="%.8f")
    age_raw_all.to_csv(RUN / "age_conditional_mar_per_seed_reconciled_latest.csv", index=False, float_format="%.8f")
    return comp, base, group_mar, type_all, age_all


def build_main(comp, type_all, age_all):
    doc = Document(MANUSCRIPT_IN)
    replace_prefix(doc, "Incomplete test features can reduce recall",
        "Incomplete test features can reduce recall for a prespecified group more than aggregate accuracy, leaving a headline score incomplete. We define the Focus-Group Masking Gap (FMG) as the signed difference between focus-group recall loss and aggregate-accuracy loss for the same frozen classifier under test-only feature masking. In the primary benchmark, ten repeated splits across three datasets and four models produced mean FMG above the study-specific descriptive reference of 0.01 in nine of 12 dataset-model comparisons at 20% MCAR. FMG is not a between-group disparity: g0 recall loss exceeded g1 in only 4/12 primary dataset-model pairs. A separate public-health extension evaluated CDC 2024 BRFSS and UCI Diabetes 130-US Hospitals data with five models, three seeds, two split settings, and three controlled masking mechanisms. At 80/20 and 20% MCAR, mean FMG ranged from 0.0272 to 0.0455 across models in BRFSS, but from 0.0000 to 0.0038 in UCI; all BRFSS seed means were positive, while none of the UCI seed means exceeded 0.01. BRFSS evaluation used survey weights and UCI splits were patient-disjoint. These results show dataset-dependent benchmark behavior, not population-level or deployment effects. FMG is a conditional diagnostic, not a fairness certificate, causal estimate, or evidence of operational utility.")
    replace_prefix(doc, "A classifier may perform well on complete data",
        "A classifier may perform well on complete data yet receive incomplete features at test time. Overall accuracy is usually the headline measure, but it can conceal a different degree of degradation for a particular demographic group. We therefore examine whether controlled feature masking at test time reduces focus-group recall more than aggregate accuracy when the trained model remains unchanged. This contrast is not itself a between-group disparity: in the primary results, g0 recall loss exceeded g1 recall loss in only 4/12 dataset-model pairs.")
    replace_prefix(doc, "Reproducibility and correction trace.",
        "Reproducibility and correction trace. The corrected primary notebook was rerun end-to-end in a CPU-based Colab runtime. The run manifest records Python 3.13.16 on Linux x86_64, two logical CPUs, and 2,032.9 seconds from primary-run start; it does not record the CPU model or memory allocation. The run produced 800 raw result rows across four datasets (German Credit exploratory), five classifiers, ten seeds, and clean, 5%, 10%, and 20% conditions. Adult adult.data contained 32,561 source rows; 2,399 rows with '?' were excluded under the complete-case rule, leaving 30,162 analyzed records. All ten CSV outputs in the run archive match the manifest's SHA-256 hashes and byte counts. Recalculation from the raw rows reproduces the twelve primary 20%-MCAR dataset-model means, standard deviations, positive-seed counts, and component losses reported in Table 3 at displayed precision; nine means exceed the study-specific descriptive reference of 0.01. The notebook's cell 33 separately regenerated the descriptive bootstrap intervals from the primary seed-level FMG values; the saved intervals match Table 3 at displayed precision. That interval CSV is separate from the ten-file full-run archive. Per-seed denominators, threshold-discordance summaries, and Adult-dependent sensitivities are retained as separate reproducibility outputs. The archive's figures are listed by hash in its manifest; the historical primary Colab run remains distinct, and its backend CPU model and resource allocation are unknown. An independent Adult rerun reported with XGBoost 3.4.1 and scikit-learn 1.8.0 produced mean XGB FMG 0.0927 (SD 0.0248), compared with 0.0890 (SD 0.0269) in the recorded run using XGBoost 3.2.0 and scikit-learn 1.6.1. Because the independent rerun also used a different hardware/thread configuration, this small discrepancy cannot be attributed uniquely to software version or hardware; exact numerical replication should use the recorded environment and settings.")
    age_lo = age_all.mean_age_mar_minus_mcar_fmg.min()
    age_hi = age_all.mean_age_mar_minus_mcar_fmg.max()
    age_max = age_all.loc[age_all.mean_age_mar_minus_mcar_fmg.idxmax()]
    age_pos_min = int(age_all.n_positive_differences.min())
    age_pos_max = int(age_all.n_positive_differences.max())
    age_text = ("Across 72 descriptive dataset-model-rate-strength contrasts, mean FMG differences (age-conditioned MAR minus matched-eligible-feature MCAR) ranged from "
        f"{age_lo:.4f} to {age_hi:.4f}; positive paired split-level differences ranged from {age_pos_min}/10 to {age_pos_max}/10. The largest positive mean difference occurred for "
        f"{age_max.dataset}/{age_max.model} at {age_max.rate_nominal:.0%} expected missingness and an age odds ratio of {age_max.age_odds_ratio_per_sd:g} per standard deviation (mean difference = {age_hi:.4f}; {int(age_max.n_positive_differences)}/10 split-level differences positive). "
        "These are descriptive repeated-split summaries, not significance tests. Effects varied across datasets, models, rates, and propensity strengths; this exploratory sensitivity does not establish a general MAR effect. Per-split values, realized masking rates, subgroup denominators, and comparison summaries are included in the supplementary analysis package.")
    replace_prefix(doc, "Across 72 descriptive dataset-model-rate-strength contrasts,", age_text)
    replace_prefix(doc, "The age-conditioned strength sweep varied by dataset,",
        f"The corrected age-conditioned sensitivity spans {age_lo:.4f} to {age_hi:.4f} across 72 tested dataset-model-rate-strength contrasts, with positive split-level differences ranging from {age_pos_min}/10 to {age_pos_max}/10. The largest positive mean difference was {age_max.dataset}/{age_max.model} at {age_max.rate_nominal:.0%} missingness and an age odds ratio of {age_max.age_odds_ratio_per_sd:g} per standard deviation ({age_hi:.4f}; {int(age_max.n_positive_differences)}/10 positive splits). This is bounded descriptive evidence for the specified mechanism and benchmark configuration, not a general MAR effect or robustness property.")
    comp_adult = comp[comp.dataset == "Adult Income"]
    g1_min, g1_max = comp_adult.mean_fmg_g1.min(), comp_adult.mean_fmg_g1.max()
    replace_prefix(doc, "At 20% MCAR, the complementary-group mean FMG was positive",
        "At 20% MCAR, corrected Adult complementary-group (g1) mean FMG ranged from "
        f"{g1_min:.4f} to {g1_max:.4f} across models, with positive g1 FMG in all ten splits for every Adult model. The g1 mean was also positive for all four models in Taiwan Credit, again with positive values in all ten splits. In COMPAS, g1 mean FMG was negative for all four models, and positive values occurred in only one to three splits. These group-level patterns vary across datasets and do not support a general claim that the prespecified focus group is always more affected. Corrected model-level summaries are reported in Supplementary Table S1.")
    type_adult = type_all[type_all.dataset == "Adult Income"]
    fmg_min, fmg_max = type_adult.mean_fmg_g0.min(), type_adult.mean_fmg_g0.max()
    pos_min, pos_max = int(type_adult.positive_fmg_g0_seeds.min()), int(type_adult.positive_fmg_g0_seeds.max())
    replace_prefix(doc, "The expanded type-aware sensitivity analysis covers RF,",
        "The expanded type-aware sensitivity analysis covers RF, three datasets, ten fixed seeds, and one 20% MCAR condition, with the same raw test-cell mask shared across imputation strategies within each dataset and seed. Imputation/encoding pipelines are fitted on clean training data; each strategy's classifier is trained on clean transformed training data and then held fixed for clean and masked test evaluation. The strategies are median imputation for numeric features plus most-frequent categorical imputation, five-neighbour KNN numeric imputation plus most-frequent categorical imputation, and iterative numeric imputation plus most-frequent categorical imputation. The in-notebook cell 35 reports only the median/mode pipeline; KNN/mode and iterative/mode results were generated separately by the type-aware sensitivity script. Table 4 and Fig. 9 report descriptive focus-group results without significance tests. Corrected Adult Income mean FMG ranges from "
        f"{fmg_min:.6f} to {fmg_max:.6f}, with {pos_min}/10 to {pos_max}/10 positive splits; COMPAS ranges from −0.0131 to 0.0049 (3–5/10 positive splits); Taiwan Credit ranges from 0.0689 to 0.0856 (10/10 positive splits). In the corrected Adult rerun, KNN/mode mean FMG was 0.032747 under scikit-learn 1.9.1, compared with 0.031762 in the earlier scikit-learn 1.6.1 output; median/mode was 0.031763 in both. The earlier KNN and median/mode values were close but distinct, and the KNN estimate shows modest version sensitivity. These type-aware means are below the primary encoded-space RF mean FMG of 0.0865, but the pipelines differ in representation and imputation; this is not an isolated causal effect or general ranking. These dataset-dependent results do not establish broad preprocessing robustness.")

    # Refresh Table 4 from corrected Adult run and the retained non-Adult rows.
    table = doc.tables[3]
    order = {"Adult Income": 0, "COMPAS": 1, "Taiwan Credit": 2}
    strat = {"median_mode": "Median/mode", "knn_numeric_mode": "KNN/mode", "iterative_numeric_mode": "Iterative/mode"}
    type_all = type_all.sort_values(["dataset", "imputation_strategy"], key=lambda s: s.map(order if s.name == "dataset" else strat))
    rows = [["Dataset", "Strategy", "FMG mean ± SD", "+ seeds", "Acc./g0 recall loss"]]
    for row in type_all.itertuples(index=False):
        rows.append([row.dataset, strat[row.imputation_strategy], f"{row.mean_fmg_g0:.6f} ± {row.sd_fmg_g0:.6f}",
                     f"{int(row.positive_fmg_g0_seeds)}/10", f"{row.mean_accuracy_drop:.4f}; {row.mean_recall_drop_g0:.4f}"])
    write_table(table, rows)
    replace_figure(doc, "Fig. 9.", RUN / "corrected_figures" / "fig09_type_aware_imputation_corrected_adult.png")
    replace_prefix(doc, "Table 4. Multi-seed, type-aware imputation sensitivity",
        "Table 4. Multi-seed, type-aware imputation sensitivity at 20% MCAR (RF; ten seeds). Adult rows were rerun with scikit-learn 1.9.1; COMPAS and Taiwan Credit retain the earlier recorded scikit-learn 1.6.1 run. Cross-dataset differences may therefore reflect both dataset and software-version differences.")
    doc.core_properties.comments = "Adult-dependent age-MAR, complementary-group, baseline-recall, and type-aware results reconciled to corrected Adult input; notebook cell scopes clarified."
    doc.save(MANUSCRIPT_OUT)


def build_supplement(comp, base, group_mar, type_all, age_all):
    doc = Document(SUPP_IN)
    comp = comp.sort_values(["dataset", "model"])
    s1 = [["Dataset", "Model", "Mean FMG g0", "Mean FMG g1", "Positive g1 splits"]]
    for r in comp.itertuples(index=False):
        s1.append([r.dataset, r.model, r4(r.mean_fmg_g0), r4(r.mean_fmg_g1), f"{int(r.positive_fmg_g1_seeds)}/10"])
    write_table(doc.tables[0], s1)

    base = base.sort_values(["dataset", "model"])
    s2 = [["Dataset", "Model", "Split pairs", "Baseline recall range", "Spearman rho"]]
    for r in base.itertuples(index=False):
        s2.append([r.dataset, r.model, str(int(r.n_seed_pairs)), f"{r.baseline_recall_min:.4f}–{r.baseline_recall_max:.4f}", f"{r.spearman_rho_baseline_recall_vs_fmg:.3f}"])
    write_table(doc.tables[1], s2)

    group_mar = group_mar.sort_values(["dataset", "model", "rate_nominal"])
    s3 = [["Dataset", "Model", "Rate", "Mean MCAR FMG", "Mean MAR FMG", "Mean difference", "Positive splits"]]
    for r in group_mar.itertuples(index=False):
        s3.append([r.dataset, r.model, f"{r.rate_nominal:.0%}", r4(r.mean_mg_mcar), r4(r.mean_mg_mar), r4(r.mean_mar_minus_mcar_mg), f"{int(r.positive_split_differences)}/10"])
    write_table(doc.tables[2], s3)

    age_all = age_all.sort_values(["dataset", "model", "rate_nominal", "age_odds_ratio_per_sd"])
    s4 = [["Dataset", "Model", "Rate", "Age OR/SD", "Mean FMG difference", "Positive splits"]]
    for r in age_all.itertuples(index=False):
        s4.append([r.dataset, r.model, f"{r.rate_nominal:.0%}", f"{r.age_odds_ratio_per_sd:g}", r4(r.mean_age_mar_minus_mcar_fmg), f"{int(r.n_positive_differences)}/10"])
    write_table(doc.tables[3], s4)

    type_all = type_all.sort_values(["dataset", "imputation_strategy"], key=lambda s: s.map({"Adult Income": 0, "COMPAS": 1, "Taiwan Credit": 2} if s.name == "dataset" else {"median_mode": 0, "knn_numeric_mode": 1, "iterative_numeric_mode": 2}))
    s5 = [["Dataset", "Pipeline", "Mean FMG ± SD", "Positive splits", "Accuracy loss", "g0 recall loss"]]
    labels = {"median_mode": "Median/mode", "knn_numeric_mode": "KNN/mode", "iterative_numeric_mode": "Iterative/mode"}
    for r in type_all.itertuples(index=False):
        s5.append([r.dataset, labels[r.imputation_strategy], f"{r.mean_fmg_g0:.6f} ± {r.sd_fmg_g0:.6f}", f"{int(r.positive_fmg_g0_seeds)}/10", r4(r.mean_accuracy_drop), r4(r.mean_recall_drop_g0)])
    write_table(doc.tables[4], s5)

    replace_prefix(doc, "Source file: `sensitivity/complementary_group_sensitivity_summary.csv`",
        "Corrected source files: `adult_corrected_rerun/complementary_group_summary_corrected_adult.csv` and `adult_corrected_rerun/complementary_group_per_seed_corrected_adult.csv`.")
    replace_prefix(doc, "Source file: `sensitivity/baseline_recall_sensitivity_summary.csv`",
        "Corrected source files: `adult_corrected_rerun/baseline_recall_summary_corrected_adult.csv` and `adult_corrected_rerun/baseline_recall_per_seed_corrected_adult.csv`.")
    replace_prefix(doc, "Source files: `sensitivity/group_conditional_mar_per_seed.csv`",
        "Corrected source files: `adult_corrected_rerun/group_conditional_mar_per_seed_with_corrected_adult.csv` and `adult_corrected_rerun/group_conditional_mar_descriptive_summary_with_corrected_adult.csv`. This constructed mechanism does not represent all MAR processes or MNAR.")
    replace_prefix(doc, "Source files: `sensitivity/age_mar_strength_outputs/age_conditional_mar_statistics.csv`",
        "Corrected combined files: `adult_corrected_rerun/age_conditional_mar_summary_reconciled_latest.csv` and `adult_corrected_rerun/age_conditional_mar_per_seed_reconciled_latest.csv`. Adult rows were rerun on the corrected UCI Adult archive; COMPAS and Taiwan rows retain the earlier recorded runs. This mechanism tests one observed covariate and a specified logistic propensity form; it does not establish a general MAR result.")
    replace_prefix(doc, "Source files: `sensitivity/typeaware_imputation_sensitivity_summary.csv`",
        "Corrected combined files: `adult_corrected_rerun/typeaware_imputation_summary_reconciled_latest.csv` and `adult_corrected_rerun/typeaware_imputation_per_seed_reconciled_latest.csv`. Adult rows were rerun on the corrected UCI Adult archive; COMPAS and Taiwan rows retain their recorded runs. The Adult KNN/mode mean FMG differed modestly between scikit-learn 1.6.1 (0.031762) and 1.9.1 (0.032747), so the table reports the latter run with its version noted. This comparison does not establish general preprocessing robustness across model families, rates, or datasets.")
    doc.core_properties.comments = "Supplementary Tables S1-S5 reconciled to corrected Adult sensitivity outputs."
    doc.save(SUPP_OUT)


if __name__ == "__main__":
    comp, base, group_mar, type_all, age_all = load_sensitivities()
    if not AGE_NEW.exists() or not TYPE_NEW.exists():
        raise FileNotFoundError("Run corrected Adult age-MAR and type-aware scripts before building manuscripts.")
    build_main(comp, type_all, age_all)
    build_supplement(comp, base, group_mar, type_all, age_all)
    print(MANUSCRIPT_OUT)
    print(SUPP_OUT)
